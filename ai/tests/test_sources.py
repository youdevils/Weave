"""
The SourceIndex substrate (ai.services.sources) and the structural readers
behind it (assisted.services.evidence_extraction): stable, addressable
segments with structural context -- never interpretation -- and safe
fallbacks.
"""

import io
from pathlib import Path

from django.test import SimpleTestCase

from assisted.services.evidence_extraction import bounded_blocks, extract_text

from ai.services.evidence_bundle import EvidenceBundle
from ai.services.sources import blocks_from_text, build_segments, mentions, value_occurs

INTERNATIONAL_FLYER = (Path(__file__).parent / "fixtures" / "international_rugby_flyer.pdf").read_bytes()


def segments_of(text=None, blocks=None):
    return EvidenceBundle.from_assets([{"name": "evidence", "content": text or "", "blocks": blocks}]).segments()


class TextSegmentationTests(SimpleTestCase):

    def test_pipe_tables_keep_rows_cells_and_header(self):
        segments = segments_of("Fixtures\nStage | Match | Venue\nPool A | New Zealand v Fiji | Eden Park\nPool B | Japan v Samoa | Sky Stadium\n")

        rows = [s for s in segments if s.kind == "table_row"]
        self.assertEqual([r.segment_id for r in rows], ["S1#t1.r1", "S1#t1.r2"])
        self.assertEqual(rows[0].cells, ["Pool A", "New Zealand v Fiji", "Eden Park"])
        self.assertEqual(rows[0].header, ["Stage", "Match", "Venue"])
        self.assertEqual(rows[0].heading_path, ["Fixtures"])
        # Structure travels as references to citable segments, never rendered text.
        self.assertEqual(rows[0].context(), {"segment_id": "S1#t1.r1", "source_id": "S1", "kind": "table_row",
                                             "text": "Pool A | New Zealand v Fiji | Eden Park", "within": ["S1#1", "S1#t1"]})

    def test_tab_separated_tables_markdown_headings_and_lists(self):
        segments = segments_of("# Venues\nVenue\tCity\nEden Park\tAuckland\n\n## Format\n- Pool Stage: two pools\n  of four teams\n- Final at Eden Park\n")

        kinds = [(s.kind, s.text) for s in segments]
        self.assertIn(("heading", "Venues"), kinds)
        self.assertIn(("table_row", "Eden Park | Auckland"), kinds)
        items = [s for s in segments if s.kind == "list_item"]
        self.assertEqual([i.text for i in items], ["- Pool Stage: two pools\nof four teams", "- Final at Eden Park"])
        self.assertEqual(items[0].heading_path, ["Venues", "Format"])

    def test_ids_are_stable_and_source_text_is_the_rendered_segments(self):
        text = "Intro paragraph.\n\nVenue | City\nEden Park | Auckland"
        first, second = segments_of(text), segments_of(text)

        self.assertEqual([s.segment_id for s in first], [s.segment_id for s in second])
        bundle = EvidenceBundle.from_assets([{"name": "n", "content": text}])
        self.assertIn("Eden Park | Auckland", bundle.sources[0].text)
        self.assertEqual(bundle.segments_containing("S1", "Eden Park | Auckland")[0].segment_id, "S1#t1.r1")

    def test_flattened_layouts_fall_back_to_one_block_per_line(self):
        blocks = blocks_from_text("Venue\nLocation\nEden Park\nAuckland", line_blocks=True)

        self.assertEqual([b["kind"] for b in blocks], ["paragraph"] * 4)

    def test_preparation_never_interprets(self):
        segments = segments_of("Pool A | New Zealand v Fiji | Eden Park\nPool B | Japan v Samoa | Sky Stadium")

        # Structure only: a two-row pipe run is a table; nothing is typed or related.
        self.assertTrue(all(set(s.model_dump()) <= {"segment_id", "source_id", "kind", "text", "heading_path", "ancestor_ids",
                                                     "header", "cells", "parent_id", "position"} for s in segments))


class PdfStructureTests(SimpleTestCase):

    def setUp(self):
        extracted = extract_text(INTERNATIONAL_FLYER, filename="flyer.pdf")
        self.segments = segments_of(extracted.text, extracted.blocks)

    def by_text(self, text):
        return next(s for s in self.segments if s.text == text)

    def test_table_rows_are_recovered_from_text_positions(self):
        row = self.by_text("Pool A | New Zealand v Fiji | Eden Park")

        self.assertEqual(row.kind, "table_row")
        self.assertEqual(row.header, ["Stage", "Match", "Venue"])
        self.assertEqual(row.heading_path[-1], "Example fixtures")

    def test_a_bold_first_row_is_a_header_and_a_plain_grid_has_none(self):
        venue = self.by_text("Eden Park | Auckland, New Zealand | Opening match and final")
        teams = self.by_text("Pool A | New Zealand | Fiji | Japan | Samoa")

        self.assertEqual(venue.header, ["Venue", "Location", "Primary use"])
        self.assertEqual(teams.header, [])
        self.assertEqual(teams.heading_path[-1], "Teams")

    def test_a_table_continued_across_pages_keeps_its_header(self):
        row = self.by_text("Quarter-final | Pool A winner v Pool B runner-up | Orangetheory Stadium")

        self.assertEqual(row.header, ["Stage", "Match", "Venue"])

    def test_headings_and_list_items_nest(self):
        final = next(s for s in self.segments if s.kind == "list_item" and s.text.startswith("- Final:"))

        self.assertEqual(final.heading_path[-1], "Competition format")
        self.assertEqual(self.segments[0].text, "PACIFIC INTERNATIONAL RUGBY CHAMPIONSHIP 2027")

    def test_bounding_keeps_whole_blocks_and_marks_truncation(self):
        extracted = extract_text(INTERNATIONAL_FLYER, filename="flyer.pdf")

        blocks, truncated = bounded_blocks(extracted.blocks, max_chars=400)

        self.assertTrue(truncated)
        self.assertIn("evidence truncated", blocks[-1]["text"])


class DocxStructureTests(SimpleTestCase):

    def test_docx_tables_headings_and_lists_are_kept(self):
        from docx import Document

        document = Document()
        document.add_heading("Venues", level=1)
        document.add_paragraph("All venues are in New Zealand.")
        table = document.add_table(rows=3, cols=2)
        for row, values in zip(table.rows, (("Venue", "City"), ("Eden Park", "Auckland"), ("Sky Stadium", "Wellington"))):
            for cell, value in zip(row.cells, values):
                cell.text = value
        document.add_paragraph("Final at Eden Park", style="List Bullet")
        buffer = io.BytesIO()
        document.save(buffer)

        extracted = extract_text(buffer.getvalue(), filename="venues.docx")
        segments = segments_of(extracted.text, extracted.blocks)

        rows = [s for s in segments if s.kind == "table_row"]
        self.assertEqual([r.text for r in rows], ["Eden Park | Auckland", "Sky Stadium | Wellington"])
        self.assertEqual(rows[0].header, ["Venue", "City"])
        self.assertEqual(rows[0].heading_path, ["Venues"])
        self.assertIn(("list_item", "Final at Eden Park"), [(s.kind, s.text) for s in segments])
        self.assertIn("Eden Park", extracted.text)


class MentionTests(SimpleTestCase):

    def test_mentions_are_whole_word_case_punctuation_and_plural_insensitive(self):
        self.assertTrue(mentions("Quarter-final | Pool A winner", ["Quarter-finals"]))
        self.assertTrue(mentions("the final played in Auckland", ["Final"]))
        self.assertFalse(mentions("Finalists: 2", ["Final"]))
        self.assertFalse(mentions("Eden Parkway", ["Eden Park"]))

    def test_values_match_numerically(self):
        self.assertTrue(value_occurs("50000", "capacity: 50,000"))
        self.assertTrue(value_occurs("Auckland", "Eden Park | Auckland, New Zealand"))
        self.assertFalse(value_occurs("48000", "capacity: 50,000"))

    def test_build_segments_assigns_heading_ancestors(self):
        segments = build_segments("S1", [{"kind": "heading", "text": "A", "level": 1}, {"kind": "heading", "text": "B", "level": 2},
                                          {"kind": "paragraph", "text": "x"}, {"kind": "heading", "text": "C", "level": 2},
                                          {"kind": "paragraph", "text": "y"}])

        self.assertEqual(segments[2].heading_path, ["A", "B"])
        self.assertEqual(segments[4].heading_path, ["A", "C"])
        self.assertEqual(segments[4].ancestor_ids, ["S1#1", "S1#4"])
