"""
The Document Evidence Model (ai.services.document) -- Phase 1 of
.Documentation/reconcile-architecture-plan.md. Purely DEM: no Reading, no
expansion, no catalogue.

Invariant 1 (the DEM is structural, never semantic), enforced:
- import boundary: the package never imports the semantic index, catalogue,
  mapping, the EvidenceGraph or anything of Reconcile's interpretation path;
- structural-only schema: no DEM type carries a field outside the allowlist,
  and the allowlist names nothing ontological;
- catalogue independence: building takes the segmented evidence and nothing
  else, and the same document always yields the same digest;
- deterministic ids: column <table>.c<k>, cell <row>.c<k>;
- structural predicates (`structurally_conforms`,
  `structured_source_segments`) need no Reading.
"""

import ast
import dataclasses
import inspect
import io
from pathlib import Path

from django.test import SimpleTestCase

from ai.services.document import build as dem_build
from ai.services.document import model as dem_model
from ai.services.document import queries
from ai.services.document.model import FIELD_ALLOWLIST, Cell, Row
from ai.services.evidence_bundle import EvidenceBundle
from ai.tests.rugby import FLYER_ASSETS
from ai.tests.test_evidence_lifecycle import flyer_assets

PACKAGE = Path(dem_model.__file__).parent

FORBIDDEN_IMPORTS = (
    "ai.services.semantic", "ai.services.reconcile", "ai.services.stages", "ai.services.evidence_graph",
    "ai.services.grounding", "ai.services.change_set", "ai.services.resolution", "model.",
)
ONTOLOGICAL_WORDS = ("type", "kind", "relationship", "hint", "role", "option", "entity", "predicate", "mapping", "catalogue")


def flyer_document():
    return EvidenceBundle.from_assets(flyer_assets()).document()


class BoundaryTests(SimpleTestCase):

    def test_the_package_imports_nothing_semantic(self):
        for path in sorted(PACKAGE.glob("*.py")):
            tree = ast.parse(path.read_text(encoding="utf-8"))
            imported = []
            for node in ast.walk(tree):
                if isinstance(node, ast.Import):
                    imported += [a.name for a in node.names]
                elif isinstance(node, ast.ImportFrom) and node.module:
                    imported.append(node.module)
            with self.subTest(module=path.name):
                self.assertFalse([m for m in imported if m.startswith(FORBIDDEN_IMPORTS)], imported)

    def test_dem_types_carry_only_structural_fields(self):
        types = [obj for _, obj in inspect.getmembers(dem_model, inspect.isclass)
                 if dataclasses.is_dataclass(obj) and obj.__module__ == dem_model.__name__]
        self.assertTrue(types)
        for cls in types:
            with self.subTest(type=cls.__name__):
                self.assertTrue({f.name for f in dataclasses.fields(cls)} <= FIELD_ALLOWLIST)
        self.assertFalse([name for name in FIELD_ALLOWLIST if any(word in name for word in ONTOLOGICAL_WORDS)])

    def test_building_depends_on_the_segmented_evidence_alone(self):
        self.assertEqual(list(inspect.signature(dem_build.build_document).parameters), ["bundle"])
        first, second = flyer_document(), flyer_document()
        self.assertEqual(first.digest, second.digest)
        self.assertEqual(first.cells, second.cells)
        self.assertNotEqual(first.digest, EvidenceBundle.from_assets(FLYER_ASSETS).document().digest)


class FlyerStructureTests(SimpleTestCase):

    def setUp(self):
        self.document = flyer_document()

    def test_ids_are_derived_from_segment_ids_and_cell_positions(self):
        self.assertEqual(self.document.columns["S1#t3.c2"].label, "Match")
        self.assertEqual(self.document.cells["S1#t3.r1.c2"].text, "New Zealand v Fiji")
        self.assertEqual(self.document.cells["S1#t3.r1.c2"].column_id, "S1#t3.c2")
        for cell in self.document.cells.values():
            self.assertEqual(cell.cell_id, f"{cell.row_id}.c{cell.ordinal}")
            self.assertEqual(cell.column_id, f"{self.document.rows[cell.row_id].table_id}.c{cell.ordinal}")

    def test_a_cell_records_text_never_meaning(self):
        # The DEM says S1#t3.r1 has Stage = Pool A, Match = New Zealand v
        # Fiji, Venue = Eden Park -- as labels and texts, nothing more.
        row = self.document.rows["S1#t3.r1"]
        self.assertEqual([(self.document.columns[c.column_id].label, c.text) for c in row.cells],
                         [("Stage", "Pool A"), ("Match", "New Zealand v Fiji"), ("Venue", "Eden Park")])
        for value in (row, *row.cells):
            self.assertEqual(set(vars(value)) <= FIELD_ALLOWLIST, True)

    def test_a_headerless_table_has_unlabelled_columns(self):
        teams = self.document.tables["S1#t1"]
        self.assertFalse(teams.has_header)
        self.assertEqual([c.column_id for c in teams.columns], [f"S1#t1.c{k}" for k in range(1, 6)])
        self.assertEqual({c.label for c in teams.columns}, {""})
        self.assertEqual(teams.ancestor_ids, ("S1#1", "S1#5"))

    def test_sections_record_containment(self):
        teams = self.document.sections["S1#5"]
        self.assertEqual((teams.heading, teams.parent_id, teams.depth), ("Teams", "S1#1", 1))
        self.assertEqual(teams.member_ids, ("S1#t1",))
        self.assertEqual(teams.descendant_ids, ("S1#t1", "S1#t1.r1", "S1#t1.r2"))
        self.assertIn("S1#5", self.document.sections["S1#1"].member_ids)

    def test_recurring_separators_are_textual_facts(self):
        signature = self.document.tables["S1#t3"].signature
        self.assertEqual(signature.separator("S1#t3.c2"), " v ")
        self.assertIsNone(signature.separator("S1#t3.c1"))
        pattern = next(p for p in self.document.tables["S1#t3"].separator_patterns if p.column_id == "S1#t3.c2" and p.separator == " v ")
        self.assertEqual((pattern.cells_with, pattern.cells_nonempty, pattern.part_counts), (4, 4, (2,)))
        # The quarter-final fixtures share the shape: whether their parts are
        # teams is a Reading's question, not the DEM's.
        self.assertEqual(self.document.tables["S1#t4"].signature.separator("S1#t4.c2"), " v ")

    def test_every_recovered_row_is_structurally_typical(self):
        for table in self.document.tables.values():
            for row in table.rows:
                with self.subTest(row=row.row_id):
                    self.assertTrue(queries.structurally_conforms(row, table.signature).conforms)

    def test_structured_segments_are_exactly_the_tables_and_rows(self):
        bundle = EvidenceBundle.from_assets(flyer_assets())
        expected = {s.segment_id for s in bundle.segments() if s.kind in ("table", "table_row")}
        self.assertEqual(queries.structured_source_segments(self.document), expected)
        self.assertEqual(queries.prose_segments(self.document), {s.segment_id for s in bundle.segments() if s.kind in ("text_block", "list_item")})
        self.assertEqual(queries.context_segments(self.document), {s.segment_id for s in bundle.segments() if s.kind == "heading"})

    def test_leads_are_where_layout_says_to_look(self):
        bundle = EvidenceBundle.from_assets(flyer_assets())
        found = {(lead["segment_id"], lead.get("column")) for lead in queries.leads(bundle, ["Pool A"], ["Match"])}
        self.assertIn(("S1#t3.r1", "Match"), found)
        self.assertIn(("S1#t3.r2", "Match"), found)
        # McLean Park sits in no row with a Match column: no lead.
        self.assertEqual(queries.leads(bundle, ["McLean Park"], ["Match"]), [])


class ConformanceTests(SimpleTestCase):

    def setUp(self):
        self.signature = flyer_document().tables["S1#t3"].signature

    def row(self, *texts, row_id="S1#t3.r9"):
        return Row(row_id=row_id, table_id="S1#t3",
                   cells=tuple(Cell(cell_id=f"{row_id}.c{k}", row_id=row_id, column_id=f"S1#t3.c{k}", ordinal=k, text=t)
                               for k, t in enumerate(texts, 1)))

    def test_structural_deviations_are_named(self):
        cases = {
            ("Pool C", "Italy v Wales", "Eden Park", "extra"): ("arity",),
            ("Stage", "Match", "Venue"): ("header_repeat",),
            ("Pool C fixtures to be confirmed", "", ""): ("spanning",),
            ("Pool C", "Italy against Wales", "Eden Park"): ("separator:S1#t3.c2",),
        }
        for texts, issues in cases.items():
            with self.subTest(texts=texts):
                self.assertEqual(queries.structurally_conforms(self.row(*texts), self.signature).issues, issues)

    def test_a_typical_row_conforms_whatever_its_text_means(self):
        # A placeholder fixture is structurally typical: rejecting it is a
        # semantic exception only a Reading may declare (invariant 17).
        self.assertTrue(queries.structurally_conforms(self.row("Quarter-final", "Pool A winner v Pool B runner-up", "Sky Stadium"),
                                                      self.signature).conforms)


class OtherSourcesTests(SimpleTestCase):

    def test_pipe_tables(self):
        text = "Venues\nVenue | City\nEden Park | Auckland\nSky Stadium | Wellington\n"
        document = EvidenceBundle.from_assets([{"name": "v.txt", "content": text, "mime_type": "text/plain"}]).document()
        table = next(iter(document.tables.values()))
        self.assertEqual([c.label for c in table.columns], ["Venue", "City"])
        self.assertEqual(document.cells[f"{table.rows[1].row_id}.c2"].text, "Wellington")
        self.assertEqual(table.ancestor_ids, ("S1#1",))

    def test_docx_tables(self):
        from docx import Document

        from assisted.services.evidence_extraction import extract_text

        doc = Document()
        doc.add_heading("Venues", level=1)
        table = doc.add_table(rows=3, cols=2)
        for r, (a, b) in enumerate((("Venue", "City"), ("Eden Park", "Auckland"), ("Sky Stadium", "Wellington"))):
            table.cell(r, 0).text, table.cell(r, 1).text = a, b
        buffer = io.BytesIO()
        doc.save(buffer)
        extracted = extract_text(buffer.getvalue(), filename="venues.docx")
        document = EvidenceBundle.from_assets([{"name": "venues.docx", "content": extracted.text, "blocks": extracted.blocks}]).document()

        table = next(iter(document.tables.values()))
        self.assertEqual([c.label for c in table.columns], ["Venue", "City"])
        self.assertEqual([c.text for c in table.rows[0].cells], ["Eden Park", "Auckland"])
        self.assertTrue(all(queries.structurally_conforms(r, table.signature).conforms for r in table.rows))
