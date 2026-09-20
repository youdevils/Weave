import datetime
import io
import zipfile
from dataclasses import replace

from ingestion.services import limits
from ingestion.services.errors import SourceFileError
from ingestion.services.parsing.table import build_table

_ZIP_MAGIC = b"PK"


def parse_xlsx(data: bytes):
    """
    A workbook with exactly one worksheet. Formulas are never evaluated:
    a formula cell contributes the value Excel last saved for it, and a
    formula that was never calculated (no saved value) rejects the file, since
    reading it as "blank" would silently mean "clear this value".
    """

    if not data.startswith(_ZIP_MAGIC):
        raise SourceFileError(
            "This is not a valid .xlsx file.",
            code="bad_signature",
        )

    has_formulas = _check_archive(data)

    import openpyxl

    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:  # openpyxl surfaces assorted parser errors
        raise SourceFileError(
            "This .xlsx file could not be read.",
            code="malformed_workbook",
        ) from exc

    try:

        if len(workbook.sheetnames) != 1:
            raise SourceFileError(
                f"The workbook has {len(workbook.sheetnames)} worksheets. "
                "Data Import supports exactly one worksheet per file.",
                code="multiple_sheets",
            )

        if not workbook.worksheets:
            raise SourceFileError(
                "The workbook does not contain a worksheet.",
                code="no_worksheet",
            )

        formula_cells = _formula_cells(data) if has_formulas else set()

        formulas_used = []

        def rows():

            for row_number, row in enumerate(workbook.worksheets[0].iter_rows(), start=1):

                if row_number > limits.MAX_SCANNED_ROWS:
                    raise SourceFileError(
                        "The worksheet has too many rows. Delete unused rows and try again.",
                        code="too_many_rows",
                    )

                values = []

                for column_number, cell in enumerate(row, start=1):

                    if cell.data_type == "e":
                        raise SourceFileError(
                            f"Cell {_coordinate(row_number, column_number)} contains an "
                            f"Excel error ({cell.value}).",
                            code="error_cell",
                        )

                    value = _plain(cell.value)

                    if (row_number, column_number) in formula_cells:

                        if value is None:
                            raise SourceFileError(
                                f"Cell {_coordinate(row_number, column_number)} contains a "
                                "formula with no calculated value. Open the file in Excel, "
                                "save it, and try again.",
                                code="uncached_formula",
                            )

                        formulas_used.append(1)

                    values.append(value)

                yield row_number, values

        # build_table drains the generator, so formulas_used is complete after.
        table = build_table("xlsx", rows())

        if formulas_used:
            table = replace(
                table,
                warnings=table.warnings
                + (f"{len(formulas_used)} formula cell(s): their last saved values are used.",),
            )

        return table

    finally:
        workbook.close()


def _check_archive(data: bytes) -> bool:
    """Zip-bomb guard. Returns whether any worksheet part contains formulas."""

    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise SourceFileError("This is not a valid .xlsx file.", code="malformed_workbook") from exc

    with archive:

        infos = archive.infolist()

        if "[Content_Types].xml" not in {info.filename for info in infos}:
            raise SourceFileError("This is not a valid .xlsx file.", code="malformed_workbook")

        total = sum(info.file_size for info in infos)

        if total > limits.xlsx_max_uncompressed_bytes():
            raise SourceFileError(
                "The workbook is too large once unpacked.",
                code="file_too_large",
            )

        ratio = limits.xlsx_max_compression_ratio()

        for info in infos:
            if info.compress_size and info.file_size / info.compress_size > ratio:
                raise SourceFileError(
                    "The workbook has an implausible compression ratio.",
                    code="file_too_large",
                )

        for info in infos:
            if _is_sheet_part(info.filename):
                blob = archive.read(info.filename)
                if b"<f>" in blob or b"<f " in blob:
                    return True

    return False


def _formula_cells(data: bytes) -> set:
    """(row, column) of every formula cell, from a second, formula-aware read."""

    import openpyxl

    found = set()

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)

    try:
        for row_number, row in enumerate(workbook.worksheets[0].iter_rows(), start=1):

            if row_number > limits.MAX_SCANNED_ROWS:
                break

            for column_number, cell in enumerate(row, start=1):
                if cell.data_type == "f":
                    found.add((row_number, column_number))
    finally:
        workbook.close()

    return found


def _is_sheet_part(name: str) -> bool:
    return name.startswith("xl/worksheets/") and name.endswith(".xml")


def _plain(value):
    """Cell values as plain data; durations and times become text."""

    if isinstance(value, datetime.timedelta):
        return str(value)

    if isinstance(value, datetime.time):
        return value.isoformat()

    return value


def _coordinate(row: int, column: int) -> str:
    letters = ""

    while column:
        column, remainder = divmod(column - 1, 26)
        letters = chr(65 + remainder) + letters

    return f"{letters}{row}"
