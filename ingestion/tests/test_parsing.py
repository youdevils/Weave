import datetime
import io
import zipfile
from pathlib import Path

import openpyxl
from django.test import SimpleTestCase, override_settings

from ingestion.services.errors import SourceFileError
from ingestion.services.parsing import parse_table

FIXTURES = Path(__file__).parent / "fixtures"


def xlsx_bytes(rows, *, extra_sheet=False, formulas=None, error_cell=None):
    workbook = openpyxl.Workbook()
    sheet = workbook.active

    for row in rows:
        sheet.append(row)

    for coordinate, formula in (formulas or {}).items():
        sheet[coordinate] = formula

    if error_cell:
        sheet[error_cell] = "#DIV/0!"
        sheet[error_cell].data_type = "e"

    if extra_sheet:
        workbook.create_sheet("Second")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def codes(callable_):
    try:
        callable_()
    except SourceFileError as error:
        return error.code
    return None


class CsvParsingTests(SimpleTestCase):

    def parse(self, text, *, encoding="utf-8"):
        return parse_table("csv", text.encode(encoding) if isinstance(text, str) else text)

    def test_reads_a_header_and_rows(self):
        table = self.parse("ID,Name\r\n1,Ann\r\n2,Bo\r\n")

        self.assertEqual(table.headers, ("ID", "Name"))
        self.assertEqual(table.rows, (("1", "Ann"), ("2", "Bo")))
        self.assertEqual(table.row_count, 2)

    def test_a_utf8_bom_is_stripped(self):
        table = parse_table("csv", b"\xef\xbb\xbfID,Name\n1,Ann\n")

        self.assertEqual(table.headers, ("ID", "Name"))

    def test_quoted_values_commas_quotes_and_newlines(self):
        table = self.parse('ID,Note\n1,"a, b"\n2,"say ""hi"""\n3,"line1\nline2"\n')

        self.assertEqual(table.rows[0][1], "a, b")
        self.assertEqual(table.rows[1][1], 'say "hi"')
        self.assertEqual(table.rows[2][1], "line1\nline2")

    def test_row_numbers_follow_the_spreadsheet_including_blank_rows(self):
        table = self.parse("ID\n1\n\n2\n")

        self.assertEqual(table.row_numbers, (2, 4))

    def test_blank_rows_and_trailing_blank_cells_are_skipped(self):
        table = self.parse("A,B,C\n\n1,2,\n,,\n3,,\n")

        self.assertEqual(table.rows, (("1", "2", None), ("3", None, None)))

    def test_empty_cells_are_none_but_interior_blanks_keep_their_position(self):
        table = self.parse("A,B,C\n1,,3\n")

        self.assertEqual(table.rows[0], ("1", None, "3"))

    def test_leading_blank_rows_before_the_header_are_ignored(self):
        table = self.parse("\n\nID\n1\n")

        self.assertEqual(table.headers, ("ID",))
        self.assertEqual(table.header_row_number, 3)

    def test_blank_headers_get_placeholders_and_duplicates_are_kept(self):
        table = self.parse("A,,A\n1,2,3\n")

        self.assertEqual(table.headers, ("A", "Column 2", "A"))

    def test_short_rows_are_padded(self):
        table = self.parse("A,B,C\n1\n")

        self.assertEqual(table.rows[0], ("1", None, None))

    def test_a_row_wider_than_the_header_is_rejected_not_truncated(self):
        self.assertEqual(codes(lambda: self.parse("A\n1,2\n")), "ragged_rows")

    def test_formula_looking_text_is_kept_verbatim(self):
        table = self.parse("A\n=1+1\n+SUM(A1)\n-5\n@cmd\n")

        self.assertEqual([row[0] for row in table.rows], ["=1+1", "+SUM(A1)", "-5", "@cmd"])

    def test_non_utf8_is_rejected_with_advice(self):
        with self.assertRaises(SourceFileError) as raised:
            self.parse("Name\ncaf\xe9\n", encoding="latin-1")

        self.assertEqual(raised.exception.code, "bad_encoding")
        self.assertIn("UTF-8", raised.exception.message)

    def test_binary_content_is_not_a_csv(self):
        self.assertEqual(codes(lambda: parse_table("csv", b"PK\x03\x04\x00\x00abc")), "not_csv")

    def test_an_empty_file_is_rejected(self):
        self.assertEqual(codes(lambda: self.parse("")), "empty_file")
        self.assertEqual(codes(lambda: self.parse("\n\n,,\n")), "empty_file")

    def test_a_header_without_data_is_rejected(self):
        self.assertEqual(codes(lambda: self.parse("A,B\n")), "no_data_rows")

    def test_a_malformed_quote_is_rejected(self):
        self.assertEqual(codes(lambda: self.parse('A,B\n1,"unterminated\n')), "malformed_csv")


class LimitTests(SimpleTestCase):

    @override_settings(IMPORT_MAX_ROWS=3)
    def test_the_row_limit(self):
        self.assertEqual(len(parse_table("csv", b"A\n1\n2\n3\n").rows), 3)
        self.assertEqual(codes(lambda: parse_table("csv", b"A\n1\n2\n3\n4\n")), "too_many_rows")

    @override_settings(IMPORT_MAX_COLUMNS=2)
    def test_the_column_limit(self):
        self.assertEqual(codes(lambda: parse_table("csv", b"A,B,C\n1,2,3\n")), "too_many_columns")

    @override_settings(IMPORT_MAX_HEADER_CHARS=5)
    def test_the_header_length_limit(self):
        self.assertEqual(codes(lambda: parse_table("csv", b"TOOLONGHEADER\n1\n")), "header_too_long")

    @override_settings(IMPORT_MAX_CELL_CHARS=5)
    def test_the_cell_length_limit(self):
        self.assertEqual(codes(lambda: parse_table("csv", b"A\nabcdefghi\n")), "cell_too_long")

    @override_settings(IMPORT_MAX_ROWS=2)
    def test_the_row_limit_applies_to_workbooks_too(self):
        data = xlsx_bytes([["A"], [1], [2], [3]])

        self.assertEqual(codes(lambda: parse_table("xlsx", data)), "too_many_rows")

    @override_settings(IMPORT_XLSX_MAX_UNCOMPRESSED_BYTES=1000)
    def test_a_workbook_that_unpacks_too_large_is_rejected(self):
        self.assertEqual(codes(lambda: parse_table("xlsx", xlsx_bytes([["A"], [1]]))), "file_too_large")

    @override_settings(IMPORT_XLSX_MAX_COMPRESSION_RATIO=2)
    def test_an_implausible_compression_ratio_is_rejected(self):
        padded = xlsx_bytes([["A"]] + [["x" * 200] for _ in range(50)])

        self.assertEqual(codes(lambda: parse_table("xlsx", padded)), "file_too_large")


class XlsxParsingTests(SimpleTestCase):

    def test_reads_typed_cells(self):
        data = xlsx_bytes(
            [
                ["ID", "Name", "Cost", "Live", "When"],
                ["A1", "Payroll", 12.5, True, datetime.datetime(2024, 5, 1)],
            ]
        )

        table = parse_table("xlsx", data)

        self.assertEqual(table.headers, ("ID", "Name", "Cost", "Live", "When"))
        self.assertEqual(table.rows[0], ("A1", "Payroll", 12.5, True, datetime.datetime(2024, 5, 1)))

    def test_blank_rows_are_skipped_and_row_numbers_are_real(self):
        data = xlsx_bytes([["A"], [], [1], [None], [2]])

        table = parse_table("xlsx", data)

        self.assertEqual(table.row_numbers, (3, 5))

    def test_exactly_one_worksheet(self):
        self.assertEqual(
            codes(lambda: parse_table("xlsx", xlsx_bytes([["A"], [1]], extra_sheet=True))),
            "multiple_sheets",
        )

    def test_a_cached_formula_uses_its_saved_value_and_is_reported(self):
        # openpyxl cannot write cached values, so build the sheet XML by hand.
        data = _workbook_with_cached_formula()

        table = parse_table("xlsx", data)

        self.assertEqual(table.rows[0], (1, 2))
        self.assertTrue(any("formula" in warning for warning in table.warnings))

    def test_a_formula_with_no_saved_value_rejects_the_file_and_names_the_cell(self):
        data = xlsx_bytes([["A", "B"], [1, None]], formulas={"B2": "=A2+1"})

        with self.assertRaises(SourceFileError) as raised:
            parse_table("xlsx", data)

        self.assertEqual(raised.exception.code, "uncached_formula")
        self.assertIn("B2", raised.exception.message)

    def test_an_excel_error_cell_rejects_the_file(self):
        data = xlsx_bytes([["A"], [1]], error_cell="A2")

        with self.assertRaises(SourceFileError) as raised:
            parse_table("xlsx", data)

        self.assertEqual(raised.exception.code, "error_cell")

    def test_a_non_workbook_is_rejected(self):
        self.assertEqual(codes(lambda: parse_table("xlsx", b"not a zip")), "bad_signature")
        self.assertEqual(codes(lambda: parse_table("xlsx", b"PK\x03\x04garbage")), "malformed_workbook")

        empty_zip = io.BytesIO()
        with zipfile.ZipFile(empty_zip, "w") as archive:
            archive.writestr("hello.txt", "hi")

        self.assertEqual(codes(lambda: parse_table("xlsx", empty_zip.getvalue())), "malformed_workbook")


class XlsParsingTests(SimpleTestCase):

    def test_reads_a_legacy_workbook(self):
        table = parse_table("xls", (FIXTURES / "apps.xls").read_bytes())

        self.assertEqual(table.headers, ("App ID", "Name", "Cost", "Live", "Go live"))
        self.assertEqual(table.row_count, 2)
        self.assertEqual(table.rows[0][:2], ("APP-001", "Payroll"))
        self.assertEqual(table.rows[0][2], 1200.0)
        self.assertIs(table.rows[0][3], True)
        self.assertEqual(table.rows[0][4], datetime.datetime(2024, 5, 1))

    def test_blank_rows_are_skipped_with_real_row_numbers(self):
        table = parse_table("xls", (FIXTURES / "apps.xls").read_bytes())

        self.assertEqual(table.row_numbers, (2, 4))

    def test_exactly_one_worksheet(self):
        self.assertEqual(
            codes(lambda: parse_table("xls", (FIXTURES / "two_sheets.xls").read_bytes())),
            "multiple_sheets",
        )

    def test_a_non_xls_is_rejected(self):
        self.assertEqual(codes(lambda: parse_table("xls", b"nope")), "bad_signature")
        self.assertEqual(
            codes(lambda: parse_table("xls", bytes.fromhex("D0CF11E0A1B11AE1") + b"garbage")),
            "malformed_workbook",
        )


class DispatchTests(SimpleTestCase):

    def test_an_unsupported_format_is_rejected(self):
        self.assertEqual(codes(lambda: parse_table("ods", b"x")), "unsupported_format")


def _workbook_with_cached_formula() -> bytes:
    """A one-sheet workbook whose B2 is `=A2+1` with the value 2 saved with it."""

    sheet = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        '<row r="1"><c r="A1" t="inlineStr"><is><t>A</t></is></c><c r="B1" t="inlineStr"><is><t>B</t></is></c></row>'
        '<row r="2"><c r="A2"><v>1</v></c><c r="B2"><f>A2+1</f><v>2</v></c></row>'
        "</sheetData></worksheet>"
    )

    files = {
        "[Content_Types].xml": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            "</Types>"
        ),
        "_rels/.rels": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
            "</Relationships>"
        ),
        "xl/workbook.xml": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>'
        ),
        "xl/_rels/workbook.xml.rels": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
            "</Relationships>"
        ),
        "xl/worksheets/sheet1.xml": sheet,
    }

    buffer = io.BytesIO()

    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)

    return buffer.getvalue()
